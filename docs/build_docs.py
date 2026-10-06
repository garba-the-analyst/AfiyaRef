"""Build the AfiyaRef presentation pack: .docx (3 parts) + .pptx (pitch deck)."""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PIn, Pt as PPt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor as PRGB

OUT = Path(__file__).parent
ACCENT = RGBColor(0x00, 0x66, 0xCC)
P_ACCENT = PRGB(0x00, 0x66, 0xCC)

# ---------------------------------------------------------------- DOCX content

PARTS = [
    ("PART 1 — INTRODUCTION", [
        ("What AfiyaRef Is", [
            "AfiyaRef is a unified digital health platform for Nigeria: one backend that powers a mobile app, a WhatsApp health assistant, and a facility web portal.",
            "It serves patients, hospitals, laboratories, pharmacies, and health insurers (HMOs) from a single shared system.",
        ]),
        ("What It Does", [
            "Geo-Proximity Hospital Finder: finds the nearest verified facilities by real GPS or a shared WhatsApp location pin, filtered by service (ICU, X-Ray, Pediatrics, Emergency), 24/7 status, and NHIA acceptance.",
            "Appointments & Lab Booking: schedule doctor consultations and diagnostic tests at listed facilities.",
            "Nurse Titi (AI First Aid Assistant): step-by-step first-aid and triage guidance with strict medical safety guardrails and a mandatory emergency disclaimer.",
            "Health Information (EHR-Lite): blood group, genotype, allergies, chronic conditions, emergency contact, and active prescriptions — a medical passport that travels with the patient.",
            "Emergency Inter-Hospital Check-in: an out-of-network patient checks in with an NHIA policy number; the platform notifies the home hospital and HMO with an EHR snapshot plus an insurance-validation request.",
        ]),
        ("How It Does It", [
            "One backend, three surfaces: a single Node.js/TypeScript API and one PostgreSQL + PostGIS database serve the mobile app (REST), the WhatsApp bot (webhooks or a locally paired session), and the facility portal (web UI).",
            "One identity: a phone number is the account. The same user exists on the app and on WhatsApp with no duplicate records.",
            "Location intelligence: PostGIS spatial queries (ST_DWithin, ST_Distance) rank facilities by true distance; OSRM computes real road routes and turn-by-turn navigation.",
            "Conversational access: a Redis-backed state machine walks WhatsApp users through find-a-hospital, Nurse Titi chat, and emergency check-in flows.",
        ]),
        ("Technologies & APIs", None),  # table inserted below
        ("Technology Table", [
            ("Backend", "Node.js + TypeScript (Express), JWT auth, Redis sessions"),
            ("Database", "PostgreSQL + PostGIS (spatial index, geography triggers)"),
            ("AI triage", "Ollama Cloud gpt-oss:20b (free tier); OpenAI-compatible fallback"),
            ("WhatsApp", "Meta Cloud API (production) / Baileys linked-device pairing (local testing)"),
            ("Maps & routing", "OpenStreetMap France tiles + Esri fallback; OSRM road routing"),
            ("Mobile app", "Expo React Native (GPS, embedded maps, in-app navigation)"),
            ("Python app", "Flet (Android APK + desktop), same backend and account"),
            ("SMS alerts", "Termii (Nigeria), with dry-run logging"),
            ("Infra/CI", "Docker Compose, Caddy + TLS, GitHub Actions APK pipeline"),
        ]),
    ]),
    ("PART 2 — TECHNICAL DETAILS", [
        ("System Architecture", [
            "Clients: Expo mobile app, Flet Android/desktop app, WhatsApp (Meta webhooks or Baileys session), facility portal (static web UI).",
            "API layer: Express routers for auth, facilities, bookings, health profiles, transfers, Nurse Titi, and WhatsApp; JWT middleware for user routes, shared-secret admin key for facility/HMO routes; rate limits on auth, chat, and webhooks.",
            "Services: auth (bcrypt + phone normalization), facility spatial search, Nurse Titi LLM with retry and offline fallback, transfer/EHR snapshot builder, Termii SMS notifier, OSRM routing proxy.",
            "Data: PostgreSQL + PostGIS; Redis holds WhatsApp conversation state, Nurse chat history, and rate-limit counters.",
        ]),
        ("Data Model", [
            "User: UUID id, unique phone_number, password hash, name, date of birth, gender, HMO/insurance provider, NHIA policy number, home hospital reference.",
            "HealthProfile (1:1): blood group, genotype, allergy and chronic-condition arrays, emergency contact, active prescriptions (JSON).",
            "Facility: name, address, phone, type (hospital/lab/pharmacy/clinic), latitude/longitude mirrored into a PostGIS geography column by trigger, services array (GIN-indexed), 24/7, NHIA, and emergency-readiness flags.",
            "Booking: user, facility, type (doctor/lab), scheduled time, status lifecycle.",
            "CrossFacilityTransfer: patient, treating and home facilities, NHIA number used, status (notified/acknowledged/rejected), full EHR + insurance payload (JSON).",
        ]),
        ("Key API Endpoints", [
            "POST /api/auth/register and /api/auth/login — phone + password identity shared across app and WhatsApp.",
            "GET /api/facilities/search?lat=&lng=&radius_km=&service= — PostGIS proximity search.",
            "GET /api/facilities/route?from_lat=&from_lng=&to_lat=&to_lng= — OSRM road geometry, distance, duration, turn-by-turn steps.",
            "POST /api/bookings, GET /api/bookings/mine — scheduling.",
            "GET/PUT /api/health-profile/me — the EHR-Lite passport.",
            "POST /api/transfers/checkin — emergency portability record + alerts.",
            "GET /api/nurse/status, POST /api/nurse/chat — triage with server-side history and emergency escalation.",
            "GET/POST /whatsapp/webhook — Meta verification handshake and message ingress with signature verification.",
            "GET /api/admin/transfers/incoming, PATCH /api/admin/transfers/:id/status — facility portal operations.",
        ]),
        ("WhatsApp Engine", [
            "A Redis state machine keyed by phone number drives IDLE, NURSE_TITI, LOCATION_AWAIT, and INTER_HOSPITAL_CHECKIN states, with a minimalist two-option mode (finder + Nurse Titi) selectable per deployment.",
            "Location attachments trigger a PostGIS top-3 search; the bot replies with a text list plus native in-chat map bubbles (no external apps).",
            "Nurse Titi mode routes every message to the LLM until EXIT; history is capped and stored per user.",
            "Two transports share one sender interface: Meta Cloud API for production, Baileys linked-device pairing for local testing without Meta onboarding.",
        ]),
        ("Nurse Titi Safety Design", [
            "System prompt constrains the model to step-by-step first aid in plain language, forbids complex diagnosis and controlled-medication prescriptions, and mandates an emergency disclaimer.",
            "The server appends the exact disclaimer to any reply missing it, so the guardrail holds regardless of model phrasing (models often use non-standard hyphens).",
            "Keyword-based emergency detection escalates with the nearest emergency-ready facilities when coordinates are available; three-attempt retry with backoff, then a safe offline fallback.",
        ]),
        ("Quality & Operations", [
            "23 automated tests (vitest + supertest) cover auth, spatial search ordering, bookings, transfers, alerts, Nurse behavior, and webhook flows; a scripted simulator drives the full bot conversation end to end.",
            "Secrets live in environment files only; admin routes require a shared key; outbound SMS runs in dry-run mode until Termii is configured.",
            "Deployment: production Compose stack (private database, Caddy TLS), documented Meta webhook cutover, mobile production URL switch.",
        ]),
    ]),
    ("PART 3 — ENTREPRENEURSHIP & FUTURE", [
        ("The Problem", [
            "Finding the right hospital in an emergency is slow and word-of-mouth driven; NHIA insurance often fails outside a patient's home hospital; first-aid knowledge gaps cost lives; patient records live on paper in single facilities.",
        ]),
        ("Customers & Market", [
            "Patients and families (WhatsApp-first access, zero install); hospitals, clinics, labs, and pharmacies (visibility, bookings, referrals); HMOs and the NHIA ecosystem (portability validation, reduced fraud); employers and schools (workforce health coverage).",
            "Beachhead: Lagos — dense facilities, high WhatsApp penetration, strong NHIA reform momentum — then other states.",
        ]),
        ("Business Model", [
            "Facility SaaS: monthly listing, booking-management, and transfer-inbox subscriptions for hospitals, labs, and pharmacies.",
            "Transaction share: commission on paid lab and consultation bookings made through the platform.",
            "HMO/NHIA integration fees: portability validation and EHR-exchange services for insurers.",
            "Future API licensing: verified facility directory and triage endpoints for third-party health apps.",
        ]),
        ("Go-To-Market", [
            "Launch the WhatsApp assistant first (no download friction) through hospital partners and community health outreaches.",
            "Seed the directory with verified Lagos facilities, then invite them to claim profiles and accept bookings.",
            "Pilot emergency check-in with 2–3 hospitals and one HMO; publish anonymized outcomes as proof.",
        ]),
        ("Competitive Moat", [
            "Dual-interface reach (app + WhatsApp + USSD-ready design) competitors rarely match.",
            "Live spatial + routing + insurance-portability loop, not just a static directory.",
            "Safety-engineered AI triage with enforced disclaimers, auditable transfer records, and real facility relationships.",
        ]),
        ("Roadmap", [
            "Phase 1 (now): Lagos directory, bot + apps, portal, SMS alerts, CI-built APK.",
            "Phase 2: in-chat booking and payments, HMO verification integrations, Yoruba/Hausa/Igbo language support, push notifications.",
            "Phase 3: multi-state expansion, self-hosted OSRM, analytics for facilities and insurers, telemedicine consultations.",
        ]),
        ("Risks & Mitigations", [
            "Clinical safety: guardrails, disclaimers, escalation, and human-in-the-loop transfers; never a diagnosis engine.",
            "Data protection: NDPR-aligned consent, minimal data collection, encrypted transport, facility-level access controls.",
            "Platform dependence: dual WhatsApp transports (Meta + paired session patterns) and an app-first experience reduce single-vendor risk.",
        ]),
        ("The Ask", [
            "We are seeking pilot hospitals, one HMO design partner, and pre-seed support to fund the Lagos launch, NHIA integration work, and the Phase 2 build.",
        ]),
    ]),
]

TABLE_MARKER = "Technology Table"


def build_docx(path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    title = doc.add_heading("AfiyaRef", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Unified Health Directory, Referral, Triage & Information Platform")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = doc.add_paragraph("Presentation pack — Parts 1–3: Introduction, Technical Details, Entrepreneurship & Future")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for part_title, sections in PARTS:
        doc.add_page_break()
        h = doc.add_heading(part_title, level=1)
        for r in h.runs:
            r.font.color.rgb = ACCENT
        for sec_title, body in sections:
            if sec_title == TABLE_MARKER:
                table = doc.add_table(rows=1, cols=2)
                table.style = "Light Grid Accent 1"
                hdr = table.rows[0].cells
                hdr[0].text = "Layer"
                hdr[1].text = "Choice"
                for layer, choice in body:  # type: ignore
                    row = table.add_row().cells
                    row[0].text = layer
                    row[1].text = choice
                doc.add_paragraph("")
                continue
            doc.add_heading(sec_title, level=2)
            for para in body or []:
                doc.add_paragraph(para, style="List Bullet")
    doc.save(path)


# ---------------------------------------------------------------- PPTX content

SLIDES = [
    ("title", "AfiyaRef", "One health platform for Nigeria — app, WhatsApp & hospitals, connected."),
    ("agenda", "Today", ["1. Introduction — what AfiyaRef is", "2. Technical deep-dive — how it works", "3. Business & future — why it wins"]),
    ("section", "Part 1 — Introduction", ""),
    ("bullets", "What AfiyaRef Is", [
        "Unified health directory, referral, triage & information platform",
        "One backend → mobile app + WhatsApp assistant + facility portal",
        "One identity: your phone number works everywhere",
    ]),
    ("bullets", "What It Does", [
        "🏥 Nearby hospital finder (GPS or WhatsApp location)",
        "📅 Doctor & lab-test booking",
        "👩🏾‍⚕️ Nurse Titi — AI first-aid & triage with safety guardrails",
        "🩺 EHR-Lite medical passport + 🚨 NHIA emergency check-in",
    ]),
    ("bullets", "How It Works", [
        "PostGIS spatial search ranks hospitals by true distance",
        "Redis state machine guides WhatsApp conversations",
        "OSRM routing + turn-by-turn navigation built in",
        "Ollama Cloud AI (free tier) with enforced safety disclaimer",
    ]),
    ("section", "Part 2 — Technical Details", ""),
    ("bullets", "Architecture", [
        "Node.js + TypeScript API · PostgreSQL + PostGIS · Redis",
        "JWT auth (app) + phone binding (WhatsApp) + admin keys (facilities)",
        "Meta Cloud API in production · Baileys pairing for local testing",
        "Expo + Flet clients share one account and one database",
    ]),
    ("bullets", "Data & Intelligence", [
        "Users, EHR profiles, geo-indexed facilities, bookings, transfers",
        "ST_DWithin / ST_Distance proximity queries",
        "Keyword emergency detection → nearest emergency-ready facilities",
        "23 automated tests + scripted end-to-end bot simulator",
    ]),
    ("bullets", "AI Safety by Design", [
        "First aid only — no diagnosis, no controlled prescriptions",
        "Disclaimer auto-appended when the model omits it",
        "Retry with backoff, then safe offline fallback",
        "Full transfer audit trail for hospitals & HMOs",
    ]),
    ("section", "Part 3 — Business & Future", ""),
    ("bullets", "Market & Model", [
        "Patients (WhatsApp-first) · hospitals/labs (SaaS) · HMOs (portability fees)",
        "Beachhead: Lagos, then nationwide",
        "Revenue: subscriptions + booking commissions + insurer integrations",
    ]),
    ("bullets", "Go-To-Market", [
        "WhatsApp-first launch — zero installs, hospital partners",
        "Pilot emergency check-in with 2–3 hospitals + 1 HMO",
        "Proof points → paid facility SaaS rollout",
    ]),
    ("bullets", "Roadmap", [
        "Now: directory, bot, apps, portal, alerts, CI-built APK",
        "Next: in-chat booking & payments, HMO verification, local languages",
        "Later: multi-state scale, telemedicine, insurer analytics",
    ]),
    ("bullets", "Risks & Mitigations", [
        "Clinical safety → guardrails + escalation + human-in-the-loop",
        "Data protection → NDPR-aligned, minimal data, encrypted transport",
        "Platform risk → dual WhatsApp transports + app-first experience",
    ]),
    ("title", "Thank You", "AfiyaRef — care, connected. Seeking pilot hospitals, one HMO partner & pre-seed support."),
]


def _add_bg(slide, prs) -> None:
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = PRGB(0x0F, 0x1B, 0x2D)


def _textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(PIn(left), PIn(top), PIn(width), PIn(height)).text_frame


def build_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = PIn(13.333)
    prs.slide_height = PIn(7.5)
    blank = prs.slide_layouts[6]

    for kind, title, body in SLIDES:
        slide = prs.slides.add_slide(blank)
        if kind == "title":
            _add_bg(slide, prs)
            tf = _textbox(slide, 1, 1.5, 11.3, 2)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = title
            r.font.size = PPt(54)
            r.font.bold = True
            r.font.color.rgb = PRGB(0xFF, 0xFF, 0xFF)
            tf2 = _textbox(slide, 1.5, 4, 10.3, 2)
            p2 = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run()
            r2.text = body
            r2.font.size = PPt(24)
            r2.font.color.rgb = PRGB(0xBB, 0xDD, 0xFF)
        elif kind == "section":
            _add_bg(slide, prs)
            tf = _textbox(slide, 1, 2.8, 11.3, 2)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = title
            r.font.size = PPt(44)
            r.font.bold = True
            r.font.color.rgb = PRGB(0xFF, 0xFF, 0xFF)
        else:
            tf = _textbox(slide, 0.8, 0.4, 11.7, 1.2)
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = title
            r.font.size = PPt(36)
            r.font.bold = True
            r.font.color.rgb = P_ACCENT
            tf2 = _textbox(slide, 1, 1.8, 11, 5)
            items = body if isinstance(body, list) else [body]
            for i, item in enumerate(items):
                para = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
                para.space_after = PPt(12)
                run = para.add_run()
                run.text = ("• " + item) if kind == "bullets" else item
                run.font.size = PPt(22)
    prs.save(path)


if __name__ == "__main__":
    build_docx(OUT / "AfiyaRef-Presentation-Pack.docx")
    build_pptx(OUT / "AfiyaRef-Pitch-Deck.pptx")
    print("docs built:")
    print(" -", OUT / "AfiyaRef-Presentation-Pack.docx")
    print(" -", OUT / "AfiyaRef-Pitch-Deck.pptx")
